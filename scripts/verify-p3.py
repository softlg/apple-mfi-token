#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""P3 验收：PPID 池 CRUD、Excel 导入逐行定位、余量预估与池实际一致、跨租户隔离。

前置：
  1. 应用以 local profile 启动（MySQL/Redis 在跑），RabbitMQ 本阶段用不到
  2. 库里有 owner 账号（P1 冒烟数据）

用法：
  PYTHONIOENCODING=utf-8 python scripts/verify-p3.py --account own_a_p2 --password Passw0rd123

维度（产品型号/SPU/PPID）每次运行带时间戳，脚本可重复执行；只在测试租户内写数据，
结束时不停清库——池子是账本，留痕便于回看，重复运行靠新维度避开"已存在"。
"""
import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request

import openpyxl

BASE = os.environ.get("MFI_BASE", "http://localhost:8090")
MYSQL = os.environ.get("MFI_MYSQL", r"C:\Program Files\MySQL\MySQL Server 8.0\bin\mysql.exe")
results = []


def check(name, passed, detail=""):
    results.append((name, bool(passed)))
    print(("  PASS  " if passed else "  FAIL  ") + name + ("" if passed else "  -> " + str(detail)[:300]))
    return passed


def ok(payload):
    return payload.get("code") == "0"


def call(method, path, body=None, token=None):
    data = None if body is None else json.dumps(body).encode("utf-8")
    request = urllib.request.Request(BASE + path, data=data, method=method)
    request.add_header("Content-Type", "application/json")
    if token:
        request.add_header("mfi-token", token)
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return response.status, json.loads(response.read().decode("utf-8") or "{}")
    except urllib.error.HTTPError as error:
        raw = error.read().decode("utf-8") or "{}"
        try:
            return error.code, json.loads(raw)
        except json.JSONDecodeError:
            return error.code, {"code": "?", "message": raw[:200], "data": None}


def download(path, token, dest):
    request = urllib.request.Request(BASE + path)
    request.add_header("mfi-token", token)
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            head = dict(response.headers)
            with open(dest, "wb") as handle:
                handle.write(response.read())
            return response.status, head
    except urllib.error.HTTPError as error:
        return error.code, {}


def upload(path, token, file_path, field="file"):
    """手搓 multipart：标准库没有 encoder，而导入接口只认 multipart/form-data"""
    boundary = "----p3verify%d" % int(time.time() * 1000)
    with open(file_path, "rb") as handle:
        content = handle.read()
    parts = ("--%s\r\nContent-Disposition: form-data; name=\"%s\"; filename=\"%s\"\r\n"
             "Content-Type: application/octet-stream\r\n\r\n" % (boundary, field, os.path.basename(file_path))
             ).encode("utf-8")
    body = parts + content + ("\r\n--%s--\r\n" % boundary).encode("utf-8")
    request = urllib.request.Request(BASE + path, data=body, method="POST")
    request.add_header("Content-Type", "multipart/form-data; boundary=" + boundary)
    request.add_header("mfi-token", token)
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            return response.status, json.loads(response.read().decode("utf-8") or "{}")
    except urllib.error.HTTPError as error:
        raw = error.read().decode("utf-8") or "{}"
        try:
            return error.code, json.loads(raw)
        except json.JSONDecodeError:
            return error.code, {"code": "?", "message": raw[:200], "data": None}


def workbook(path, rows):
    handle = openpyxl.Workbook()
    sheet = handle.active
    sheet.append(["PPID", "SPU", "产品型号", "可申请总量", "备注"])
    for row in rows:
        sheet.append(row)
    handle.save(path)
    return path


def mysql_row(sql):
    """账号密码走一次性 defaults 文件，不出现在命令行与进程列表里"""
    handle = tempfile.NamedTemporaryFile("w", suffix=".cnf", delete=False, encoding="utf-8")
    handle.write("[client]\nuser=%s\npassword=%s\n" % (
        os.environ.get("MFI_DB_USERNAME", "root"), os.environ.get("MFI_DB_PASSWORD", "root")))
    handle.close()
    try:
        proc = subprocess.run([MYSQL, "--defaults-extra-file=" + handle.name, "-N", "-B", "-e", sql,
                               "apple_mfi_token"], capture_output=True, timeout=60)
        out = proc.stdout.decode("utf-8", "replace").strip()
        if proc.returncode != 0:
            print("  mysql 调用失败: " + proc.stderr.decode("utf-8", "replace")[:200])
        return out
    finally:
        os.unlink(handle.name)


def cells(path):
    return [[c.value for c in r] for r in openpyxl.load_workbook(path).active.iter_rows()]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--account", default="own_a_p2")
    parser.add_argument("--password", default="Passw0rd123")
    parser.add_argument("--tmp", default="logs/p3-tmp")
    args = parser.parse_args()
    os.makedirs(args.tmp, exist_ok=True)

    stamp = time.strftime("%H%M%S")
    model = "P3-M-" + stamp
    spu = "P3-SPU-" + stamp

    status, login = call("POST", "/api/auth/login", {"account": args.account, "password": args.password})
    if not (status == 200 and ok(login)):
        print("登录失败，后续断言无意义: %s %s" % (status, json.dumps(login, ensure_ascii=False)[:300]))
        sys.exit(1)
    payload = login["data"]
    token = payload["tokenValue"]
    tenant_id = payload["currentTenant"]["tenantId"]
    print("已登录 tenantId=%s owner=%s，维度 %s / %s" % (tenant_id, payload["currentTenant"]["owner"], model, spu))

    print("\n[1] 导入模板")
    template_path = os.path.join(args.tmp, "template.xlsx")
    code, head = download("/api/ppid/template", token, template_path)
    check("模板可下载且是 xlsx", code == 200 and open(template_path, "rb").read(2) == b"PK", code)
    check("模板带文件名且按 RFC 5987 编码", "filename*=UTF-8''" in head.get("Content-Disposition", ""),
          head.get("Content-Disposition"))
    headers = [c.value for c in openpyxl.load_workbook(template_path).active[1]]
    check("模板表头列序固定", headers == ["PPID", "SPU", "产品型号", "可申请总量", "备注"], headers)

    print("\n[2] 手工新增与分页筛选")
    created = {}
    for ppid, total in [("P3-A-" + stamp, 100), ("P3-B-" + stamp, 30), ("P3-C-" + stamp, 5)]:
        status, resp = call("POST", "/api/ppid",
                            {"ppid": ppid, "spu": spu, "productModel": model,
                             "totalQuantity": total, "remark": "P3 验收"}, token)
        check("新增 " + ppid, ok(resp) and (resp.get("data") or {}).get("remainingQuantity") == total,
              resp.get("message"))
        created[ppid] = ((resp.get("data") or {}).get("id"), total)
    status, dup = call("POST", "/api/ppid",
                       {"ppid": "P3-A-" + stamp, "spu": spu, "productModel": model, "totalQuantity": 7}, token)
    check("PPID+SPU 重复被拒而不是覆盖", status == 400 and dup.get("code") == "A0422",
          "%s %s" % (status, dup.get("message")))
    status, bad = call("POST", "/api/ppid",
                       {"ppid": "", "spu": spu, "productModel": model, "totalQuantity": 0}, token)
    check("空 PPID 与非正数量被参数校验挡住", status == 400 and bad.get("code") == "A0400", bad.get("code"))
    _, page = call("GET", "/api/ppid/page?pageSize=50&productModel=" + model, token=token)
    rows = (page.get("data") or {}).get("records") or []
    check("按产品型号筛选命中 3 行", len(rows) == 3, len(rows))
    check("分页元信息齐全", (page.get("data") or {}).get("total") == 3
          and (page.get("data") or {}).get("pageNum") == 1, page.get("data"))
    _, like = call("GET", "/api/ppid/page?ppid=P3-B&productModel=" + model, token=token)
    check("PPID 模糊匹配只命中 1 行", len(((like.get("data") or {}).get("records") or [])) == 1,
          like.get("data"))

    print("\n[3] 余量预估与池实际一致")
    _, avail = call("GET", "/api/ppid/available?productModel=%s&spu=%s&quantity=120" % (model, spu), token=token)
    data = avail.get("data") or {}
    plan = data.get("plan") or []
    check("预估覆盖全部余量", data.get("allocatableQuantity") == 135, data.get("allocatableQuantity"))
    check("120 份够用并拆成 2 片", data.get("sufficient") is True and len(plan) == 2,
          json.dumps(data, ensure_ascii=False)[:200])
    check("拆分按余量倒序", [p["ppid"] for p in plan] == ["P3-A-" + stamp, "P3-B-" + stamp],
          [p["ppid"] for p in plan])
    check("分片数量合计等于请求量", sum(p["quantity"] for p in plan) == 120, plan)
    _, tight = call("GET", "/api/ppid/available?productModel=%s&spu=%s&quantity=136" % (model, spu), token=token)
    check("差 1 份就判不够", (tight.get("data") or {}).get("sufficient") is False, tight.get("data"))
    check("不够时计划把能给的都给完",
          sum(p["quantity"] for p in ((tight.get("data") or {}).get("plan") or [])) == 135, tight.get("data"))
    sql_pool = mysql_row("select sum(remaining_quantity) from t_ppid where tenant_id=%s and product_model='%s'"
                         " and spu='%s' and status=1 and deleted=0" % (tenant_id, model, spu))
    check("预估余量与库内实际一致", sql_pool.strip() == "135", sql_pool)

    print("\n[4] 修改规则：账本已动过的额度不能随便改")
    id_a = created["P3-A-" + stamp][0]
    id_c = created["P3-C-" + stamp][0]
    status, off = call("PUT", "/api/ppid/%s" % id_c, {"totalQuantity": 5, "status": 0, "remark": "先停"}, token)
    check("停用一行", ok(off) and (off.get("data") or {}).get("status") == 0, off.get("message"))
    _, after_off = call("GET", "/api/ppid/available?productModel=%s&spu=%s&quantity=1" % (model, spu), token=token)
    check("停用行立刻从预估里消失",
          (after_off.get("data") or {}).get("allocatableQuantity") == 130, after_off.get("data"))
    # 用 SQL 把已申请量抬起来，模拟 P4 之后"这份额度被申请占用"的账本
    mysql_row("update t_ppid set applied_quantity=20, remaining_quantity=10 where id=%s" % id_a)
    status, low = call("PUT", "/api/ppid/%s" % id_a, {"totalQuantity": 19, "status": 1, "remark": None}, token)
    check("总量不能低于已申请量", status == 400 and "20" in (low.get("message") or ""),
          "%s %s" % (status, low.get("message")))
    status, zero = call("PUT", "/api/ppid/%s" % id_a, {"totalQuantity": 20, "status": 1, "remark": None}, token)
    check("余量为 0 不允许启用", status == 400 and "余量" in (zero.get("message") or ""), zero.get("message"))
    _, keep = call("GET", "/api/ppid/page?ppid=P3-A&productModel=" + model, token=token)
    kept = ((keep.get("data") or {}).get("records") or [{}])[0]
    check("被拒的修改没有动过账本", kept.get("appliedQuantity") == 20 and kept.get("remainingQuantity") == 10,
          json.dumps(kept, ensure_ascii=False)[:200])
    status, grow = call("PUT", "/api/ppid/%s" % id_a, {"totalQuantity": 120, "status": 1, "remark": "扩容"}, token)
    check("提高总量后余量按快照重算",
          ok(grow) and (grow.get("data") or {}).get("remainingQuantity") == 100, grow.get("message"))

    print("\n[5] 停用不是删除")
    status, gone = call("DELETE", "/api/ppid/%s" % id_c, token=token)
    check("DELETE 只停用", status == 200 and ok(gone), gone.get("message"))
    still = mysql_row("select count(*), sum(remaining_quantity) from t_ppid where id=%s and deleted=0" % id_c)
    check("行与余量都留在库里对账", still.strip().startswith("1\t5"), still)
    _, only_off = call("GET", "/api/ppid/page?productModel=%s&status=0" % model, token=token)
    offs = ((only_off.get("data") or {}).get("records") or [])
    check("状态筛选能看到停用行", [r["id"] for r in offs] == [id_c], json.dumps(offs, ensure_ascii=False)[:200])

    print("\n[6] Excel 导入：逐行定位与整份留走")
    imp_model = model + "-IMP"
    imp_spu = spu + "-IMP"
    before = mysql_row("select count(*) from t_ppid where tenant_id=%s and product_model='%s'"
                       % (tenant_id, imp_model)).strip()
    bad_path = workbook(os.path.join(args.tmp, "bad.xlsx"), [
        ["P3-IMP-1", imp_spu, imp_model, 100, "合法"],
        ["P3-IMP-2", None, imp_model, 100, "SPU 空"],
        ["P3-IMP-3", imp_spu, imp_model, "12.5", "不是整数"],
        ["P3-IMP-1", imp_spu, imp_model, 50, "文件内重复"],
        ["P3-IMP-5", imp_spu, imp_model, 0, "零额度"]])
    status, res = upload("/api/ppid/import", token, bad_path)
    body = res.get("data") or {}
    check("含错文件整份不入库", ok(res) and body.get("insertedRows") == 0, res.get("message"))
    check("读到 5 行并拒 4 行", body.get("readRows") == 5 and body.get("rejectedRows") == 4,
          json.dumps(body, ensure_ascii=False)[:300])
    check("行号就是 Excel 界面行号（表头算第 1 行）", [e["rowNo"] for e in body.get("errors") or []] == [3, 4, 5, 6],
          body.get("errors"))
    reasons = " | ".join(e.get("message") or "" for e in body.get("errors") or [])
    check("每行原因可读且各不不同", "SPU" in reasons and "整数" in reasons and "重复" in reasons
          and "至少为 1" in reasons, reasons[:300])
    after_bad = mysql_row("select count(*) from t_ppid where tenant_id=%s and product_model='%s'"
                          % (tenant_id, imp_model)).strip()
    check("库里行数没变", before == after_bad == "0", "%s -> %s" % (before, after_bad))

    bulk = [[("P3-IMP-%04d" % i), imp_spu, imp_model, (100 if i % 3 else "100.0"), None]
            for i in range(1, 1201)]
    bulk_path = workbook(os.path.join(args.tmp, "bulk.xlsx"), bulk)
    status, res = upload("/api/ppid/import", token, bulk_path)
    body = res.get("data") or {}
    check("1200 行一次导入成功", ok(res) and body.get("insertedRows") == 1200,
          json.dumps(body, ensure_ascii=False)[:200])
    _, stats = call("GET", "/api/ppid/stats", token=token)
    mine = [s for s in (stats.get("data") or []) if s.get("productModel") == imp_model]
    check("总览按 型号+SPU 聚合且数值与导入一致",
          len(mine) == 1 and mine[0].get("ppidRows") == 1200 and mine[0].get("totalQuantity") == 120000
          and mine[0].get("remainingQuantity") == 120000 and mine[0].get("appliedQuantity") == 0,
          json.dumps(mine, ensure_ascii=False)[:250])
    check("被 Excel 存成数字的 100.0 也认",
          mysql_row("select sum(total_quantity) from t_ppid where tenant_id=%s and product_model='%s'"
                    % (tenant_id, imp_model)).strip() == "120000")
    status, res = upload("/api/ppid/import", token, bulk_path)
    body = res.get("data") or {}
    check("重复导入同一份文件全部按已存在拒绝",
          body.get("insertedRows") == 0 and body.get("rejectedRows") == 1200
          and "已存在" in ((body.get("errors") or [{}])[0].get("message") or ""),
          json.dumps(body, ensure_ascii=False)[:250])
    # 1199 行全新 + 1 行与库里撞：整份退回，一行都不许进库
    mixed = [[("P3-IMP-%04d" % i), imp_spu, imp_model, 10, None] for i in range(2001, 3200)]
    mixed.append(["P3-IMP-0001", imp_spu, imp_model, 10, None])
    mixed_path = workbook(os.path.join(args.tmp, "mixed.xlsx"), mixed)
    status, res = upload("/api/ppid/import", token, mixed_path)
    body = res.get("data") or {}
    check("与库冲突的那一行会让整份退回",
          body.get("insertedRows") == 0 and body.get("rejectedRows") == 1
          and (body.get("errors") or [{}])[0].get("rowNo") == 1201, json.dumps(body, ensure_ascii=False)[:250])
    check("退回后库里行数仍然是 1200 行",
          mysql_row("select count(*) from t_ppid where tenant_id=%s and product_model='%s'"
                    % (tenant_id, imp_model)).strip() == "1200")

    print("\n[7] 导出")
    export_path = os.path.join(args.tmp, "export.xlsx")
    code, head = download("/api/ppid/export?productModel=%s" % imp_model, token, export_path)
    rows = cells(export_path)
    check("导出可下载", code == 200 and rows[0][0] == "PPID", code)
    check("导出行数等于库里行数", len(rows) - 1 == 1200, len(rows) - 1)
    check("导出带计数列与状态", rows[0][4:7] == ["已申请量", "剩余可申请", "状态"], rows[0])
    check("状态是中文而不是 0/1", set(r[6] for r in rows[1:]) == {"可用"}, set(r[6] for r in rows[1:]))
    empty_path = os.path.join(args.tmp, "export-empty.xlsx")
    code, head = download("/api/ppid/export?productModel=NO-SUCH-MODEL", token, empty_path)
    check("命中 0 行也给得出只有表头的文件", code == 200 and len(cells(empty_path)) == 1,
          "%s %s" % (code, cells(empty_path)))

    print("\n[8] 跨租户隔离")
    other = "p3own_%s" % stamp
    status, reg = call("POST", "/api/auth/register",
                       {"account": other, "password": args.password, "name": "P3 隔离租户主",
                        "email": other + "@example.com", "tenantName": "P3 隔离租户 " + stamp})
    token_b = ((reg.get("data") or {}).get("tokenValue")) if ok(reg) else None
    if not token_b:
        status, again = call("POST", "/api/auth/login", {"account": other, "password": args.password})
        token_b = (again.get("data") or {}).get("tokenValue")
    tenant_b = ((reg.get("data") or {}).get("currentTenant") or {}).get("tenantId")
    if not token_b:
        check("第二租户可注册/登录", False, json.dumps(reg, ensure_ascii=False)[:200])
    else:
        check("第二租户可注册/登录", True)
        _, sees = call("GET", "/api/ppid/page?pageSize=200&productModel=" + model, token=token_b)
        check("B 租户按同型号查不到 A 的行", ((sees.get("data") or {}).get("records") or []) == [],
              sees.get("data"))
        status, steal = call("PUT", "/api/ppid/%s" % id_a,
                             {"totalQuantity": 999999, "status": 1, "remark": None}, token_b)
        check("B 租户改不动 A 的行", status == 400 and "不存在" in (steal.get("message") or ""),
              "%s %s" % (status, steal.get("message")))
        status, steal2 = call("DELETE", "/api/ppid/%s" % id_c, token=token_b)
        check("B 租户停不掉 A 的行", status == 400, "%s %s" % (status, steal2.get("message")))
        check("A 的行状态没被越权改动",
              mysql_row("select status from t_ppid where id=%s" % id_c).strip() == "0")
        _, avail_b = call("GET", "/api/ppid/available?productModel=%s&spu=%s&quantity=1" % (model, spu),
                          token=token_b)
        check("B 租户预估为 0", (avail_b.get("data") or {}).get("allocatableQuantity") == 0, avail_b.get("data"))
        same_ppid = workbook(os.path.join(args.tmp, "other.xlsx"),
                             [["P3-IMP-0001", imp_spu, imp_model, 7, "同 PPID 另一租户"]])
        status, res = upload("/api/ppid/import", token_b, same_ppid)
        body = res.get("data") or {}
        check("同 PPID+SPU 在另一租户可以导入（去重不跨租户）",
              ok(res) and body.get("insertedRows") == 1, json.dumps(body, ensure_ascii=False)[:250])
        export_b = os.path.join(args.tmp, "export-b.xlsx")
        code, head = download("/api/ppid/export?productModel=%s" % imp_model, token_b, export_b)
        check("B 租户导出只含自己的行", code == 200 and len(cells(export_b)) - 1 == 1, len(cells(export_b)) - 1)

    print("\n[9] 接口可见性")
    _, docs = call("GET", "/v3/api-docs")
    paths = json.dumps(docs)
    for path in ["/api/ppid/page", "/api/ppid", "/api/ppid/available", "/api/ppid/import",
                 "/api/ppid/template", "/api/ppid/export", "/api/ppid/stats"]:
        check("openapi 暴露 " + path, path in paths, None)
    check("openapi 不泄漏库列名", "remaining_quantity" not in paths and "create_by" not in paths, None)

    failed = [name for name, passed in results if not passed]
    print("\n==== 结果: %d/%d 通过 ====" % (len(results) - len(failed), len(results)))
    print("维度留档：tenant=%s model=%s spu=%s（另一个租户 %s / %s）"
          % (tenant_id, model, spu, other, tenant_b))
    for name in failed:
        print("  未通过: " + name)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()




